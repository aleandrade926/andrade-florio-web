import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
import json
import zipfile

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def build_docx_contract(filename, title, subtitle, orientacao, preambulo, clausulas, out_dir):
    doc = docx.Document()
    
    # Header
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_t = p_t.add_run(title.upper())
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(13)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(10, 25, 47)
    
    if subtitle:
        p_s = doc.add_paragraph()
        p_s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_s = p_s.add_run(subtitle)
        run_s.font.name = 'Times New Roman'
        run_s.font.size = Pt(10)
        run_s.font.italic = True
        run_s.font.color.rgb = RGBColor(90, 90, 90)
    
    doc.add_paragraph()

    # Orientação Jurídica da Banca
    if orientacao:
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F8F9FA")
        p = cell.paragraphs[0]
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.15)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run("ORIENTAÇÃO TÉCNICA E SEGURANÇA JURÍDICA (ANDRADE & FLORIO):\n" + orientacao)
        run.font.name = 'Arial'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(140, 100, 15)
        run.font.bold = True
        doc.add_paragraph()

    # Preâmbulo
    p_preamb = doc.add_paragraph()
    p_preamb.paragraph_format.line_spacing = 1.25
    p_preamb.paragraph_format.space_after = Pt(12)
    p_preamb.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run_p = p_preamb.add_run(preambulo)
    run_p.font.name = 'Times New Roman'
    run_p.font.size = Pt(11)
    run_p.font.color.rgb = RGBColor(30, 30, 30)

    # Cláusulas
    for c_title, c_body in clausulas:
        p_ct = doc.add_paragraph()
        p_ct.paragraph_format.space_before = Pt(12)
        p_ct.paragraph_format.space_after = Pt(3)
        run_ct = p_ct.add_run(c_title)
        run_ct.font.name = 'Times New Roman'
        run_ct.font.size = Pt(11)
        run_ct.font.bold = True
        run_ct.font.color.rgb = RGBColor(10, 25, 47)

        p_cb = doc.add_paragraph()
        p_cb.paragraph_format.space_before = Pt(0)
        p_cb.paragraph_format.space_after = Pt(8)
        p_cb.paragraph_format.line_spacing = 1.25
        p_cb.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run_cb = p_cb.add_run(c_body)
        run_cb.font.name = 'Times New Roman'
        run_cb.font.size = Pt(10.5)
        run_cb.font.color.rgb = RGBColor(30, 30, 30)

    # Fechamento e Assinaturas
    p_close = doc.add_paragraph()
    p_close.paragraph_format.space_before = Pt(18)
    p_close.paragraph_format.space_after = Pt(24)
    run_cl = p_close.add_run(
        "E, por estarem assim justas, acordadas e contratadas, as partes firmam o presente instrumento em 02 (duas) vias de igual teor e forma, na presença das 02 (duas) testemunhas instrumentárias abaixo qualificadas, que a tudo assistiram e dão fé.\n\n"
        "[CIDADE/UF], [DIA] de [MÊS] de 202[X].\n\n\n"
        "___________________________________________________________\n"
        "PRIMEIRA PARTE\n"
        "CPF/CNPJ nº [NÚMERO]\n\n\n"
        "___________________________________________________________\n"
        "SEGUNDA PARTE\n"
        "CPF/CNPJ nº [NÚMERO]\n\n\n"
        "TESTEMUNHAS INSTRUMENTÁRIAS:\n\n"
        "1. ______________________________________\n"
        "Nome: [NOME COMPLETO DA TESTEMUNHA 1]\n"
        "RG nº [NÚMERO] | CPF nº [NÚMERO]\n\n"
        "2. ______________________________________\n"
        "Nome: [NOME COMPLETO DA TESTEMUNHA 2]\n"
        "RG nº [NÚMERO] | CPF nº [NÚMERO]"
    )
    run_cl.font.name = 'Times New Roman'
    run_cl.font.size = Pt(10.5)

    fpath = os.path.join(out_dir, filename)
    doc.save(fpath)
    return fpath

print("Módulo de geração de minutas robustas carregado.")
