import os
import json
import shutil
import zipfile
import docx

ROOT = r"C:\Users\Alexandre\andrade-florio-web-recuperado"
DOWNLOADS_DIR = os.path.join(ROOT, "public", "downloads")
JSON_PATH = os.path.join(ROOT, "src", "data", "contratosData.json")

# Map of contract ID to docx filename
CONTRACT_MAP = [
    {
        "id": "contrato-01",
        "title": "01. Contrato de Compra e Venda de Imóvel (À Vista e Parcelado)",
        "badge": "Mais Buscado",
        "file": "01_Contrato_Compra_e_Venda_Imovel.docx",
        "alt_file": "01_Contrato_Compra_e_Venda_Imovel.docx"
    },
    {
        "id": "contrato-02",
        "title": "02. Contrato de Locação Residencial com Garantias",
        "badge": "Lei do Inquilinato",
        "file": "02_Contrato_Locacao_Residencial.docx",
        "alt_file": "02_Contrato_Locacao_Residencial.docx"
    },
    {
        "id": "contrato-03",
        "title": "03. Contrato de Locação Comercial com Regras de Ponto",
        "badge": "Comercial / PJ",
        "file": "03_Contrato_Locacao_Comercial.docx",
        "alt_file": "03_Contrato_Locacao_Comercial.docx"
    },
    {
        "id": "contrato-04",
        "title": "04. Contrato de Comodato de Imóvel (Anti-Usucapião)",
        "badge": "Proteção Patrimonial",
        "file": "04_Contrato_Comodato_Imovel.docx",
        "alt_file": "04_Contrato_Comodato_Imovel.docx"
    },
    {
        "id": "contrato-05",
        "title": "05. Termo de Vistoria de Imóvel (Entrada e Saída)",
        "badge": "Essencial para Locação",
        "file": "05_Termo_Vistoria_Imovel.docx",
        "alt_file": "05_Termo_Vistoria_Imovel.docx"
    },
    {
        "id": "contrato-06",
        "title": "06. Termo Aditivo de Contrato de Locação (Prorrogação/Reajuste)",
        "badge": "Prorrogação Segura",
        "file": "06_Aditivo_Contratual_Locacao.docx",
        "alt_file": "06_Aditivo_Contratual_Locacao.docx"
    },
    {
        "id": "contrato-07",
        "title": "07. Distrato de Contrato de Locação (Encerramento Amigável)",
        "badge": "Encerramento Seguro",
        "file": "07_Distrato_de_Locacao.docx",
        "alt_file": "07_Distrato_de_Locacao.docx"
    },
    {
        "id": "contrato-08",
        "title": "08. Notificação Extrajudicial para Desocupação / Cobrança de Aluguel",
        "badge": "Cobrança Formal",
        "file": "08_Notificacao_Extrajudicial_Despejo.docx",
        "alt_file": "08_Notificacao_Extrajudicial_Despejo.docx"
    },
    {
        "id": "contrato-09",
        "title": "09. Recibo de Sinal e Princípio de Pagamento (Arras de Imóvel)",
        "badge": "Arras / Compra e Venda",
        "file": "09_Recibo_Sinal_e_Arras.docx",
        "alt_file": "09_Recibo_Sinal_e_Arras.docx"
    },
    {
        "id": "contrato-10",
        "title": "10. Contrato de Corretagem Imobiliária e Intermediação",
        "badge": "Comissão Garantida",
        "file": "10_Contrato_Corretagem_Imobiliaria.docx",
        "alt_file": "10_Contrato_Corretagem_Imobiliaria.docx"
    },
    {
        "id": "contrato-11",
        "title": "11. Contrato de Empreitada e Reforma de Imóvel",
        "badge": "Obras e Reformas",
        "file": "11_Contrato_Empreitada_Reforma.docx",
        "alt_file": "11_Contrato_Empreitada_Reforma.docx"
    },
    {
        "id": "contrato-12",
        "title": "12. Contrato de Prestação de Serviços Geral",
        "badge": "Escopo e Prazos",
        "file": "12_Contrato_Prestacao_Servicos_Simples.docx",
        "alt_file": "12_Contrato_Prestacao_Servicos_Simples.docx"
    },
    {
        "id": "contrato-13",
        "title": "13. Contrato de Prestação de Serviços PJ / Freelancer (Anti-CLT)",
        "badge": "Blindagem Trabalhista",
        "file": "13_Contrato_Servicos_PJ_Freelancer.docx",
        "alt_file": "13_Contrato_Servicos_PJ_Freelancer.docx"
    },
    {
        "id": "contrato-14",
        "title": "14. Acordo de Confidencialidade e Não Divulgação (NDA)",
        "badge": "Segredo Comercial",
        "file": "14_Acordo_Confidencialidade_NDA.docx",
        "alt_file": "14_Acordo_Confidencialidade_NDA.docx"
    },
    {
        "id": "contrato-15",
        "title": "15. Contrato de Parceria Comercial e Joint Venture",
        "badge": "Divisão de Resultados",
        "file": "15_Contrato_Parceria_Comercial.docx",
        "alt_file": "15_Contrato_Parceria_Comercial.docx"
    },
    {
        "id": "contrato-16",
        "title": "16. Contrato de Representação Comercial Autônoma",
        "badge": "Lei 4.886/65",
        "file": "16_Contrato_Representacao_Comercial.docx",
        "alt_file": "16_Contrato_Representacao_Comercial.docx"
    },
    {
        "id": "contrato-17",
        "title": "17. Termo de Encerramento e Quitação Plena de Serviços",
        "badge": "Quitação Definitiva",
        "file": "17_Termo_Quitacao_Servicos.docx",
        "alt_file": "17_Termo_Quitacao_Servicos.docx"
    },
    {
        "id": "contrato-18",
        "title": "18. Distrato de Contrato de Prestação de Serviços",
        "badge": "Rescisão Amigável",
        "file": "18_Distrato_Prestacao_Servicos.docx",
        "alt_file": "18_Distrato_Prestacao_Servicos.docx"
    },
    {
        "id": "contrato-19",
        "title": "19. Contrato de Licenciamento de Uso de Marca",
        "badge": "Propriedade Intelectual",
        "file": "19_Licenciamento_Uso_Marca.docx",
        "alt_file": "19_Licenciamento_Uso_Marca.docx"
    },
    {
        "id": "contrato-20",
        "title": "20. Contrato de Cessão Total de Direitos Autorais e Imagem",
        "badge": "Lei 9.610/98",
        "file": "20_Cessao_Direitos_Autorais.docx",
        "alt_file": "20_Cessao_Direitos_Autorais.docx"
    },
    {
        "id": "contrato-21",
        "title": "21. Contrato de Compra e Venda de Veículo Usado (Com Vistoria)",
        "badge": "Art. 134 CTB",
        "file": "21_Contrato_Compra_Venda_Veiculo_Usado.docx",
        "alt_file": "21_Contrato_Compra_Venda_Veiculo.docx"
    },
    {
        "id": "contrato-22",
        "title": "22. Termo de Confissão de Dívida e Parcelamento",
        "badge": "Título Executivo (CPC)",
        "file": "22_Confissao_Divida_Parcelamento_Titulo_Executivo.docx",
        "alt_file": "22_Confissao_de_Divida.docx"
    },
    {
        "id": "contrato-23",
        "title": "23. Contrato de Mútuo Financeiro (Empréstimo de Dinheiro)",
        "badge": "Lei de Usura / CC",
        "file": "23_Contrato_Mutuo_Financeiro_Entre_Particulares.docx",
        "alt_file": "23_Contrato_Mutuo_Emprestimo.docx"
    },
    {
        "id": "contrato-24",
        "title": "24. Termo de Acordo Extrajudicial e Transação",
        "badge": "Quitação Ampla",
        "file": "24_Termo_Acordo_Extrajudicial_Transacao_Quitacao.docx",
        "alt_file": "24_Termo_Acordo_Extrajudicial.docx"
    },
    {
        "id": "contrato-25",
        "title": "25. Notificação Extrajudicial de Cobrança com Prazo",
        "badge": "Constituição em Mora",
        "file": "25_Notificacao_Extrajudicial_Cobranca_Constituicao_Mora.docx",
        "alt_file": "25_Notificacao_Extrajudicial_Cobranca.docx"
    },
    {
        "id": "contrato-26",
        "title": "26. Procuração de Plenos Poderes (Ad Negotia)",
        "badge": "Gestão e Bancos",
        "file": "26_Procuracao_Plenos_Poderes_Extrajudicial_Ad_Negotia.docx",
        "alt_file": "26_Procuracao_Plenos_Poderes.docx"
    },
    {
        "id": "contrato-27",
        "title": "27. Procuração com Poderes Específicos para Venda de Imóvel",
        "badge": "Art. 661, §1º CC",
        "file": "27_Procuracao_Poderes_Especificos_Venda_Imovel.docx",
        "alt_file": "27_Procuracao_Especifica_Venda.docx"
    },
    {
        "id": "contrato-28",
        "title": "28. Declaração de Hipossuficiência Financeira (Justiça Gratuita)",
        "badge": "Art. 98 CPC",
        "file": "28_Declaracao_Hipossuficiencia_Gratuidade_Justica.docx",
        "alt_file": "28_Declaracao_Hipossuficiencia.docx"
    },
    {
        "id": "contrato-29",
        "title": "29. Contrato de Doação de Bens com Reserva de Usufruto",
        "badge": "Proteção Familiar",
        "file": "29_Contrato_Doacao_Imovel_Reserva_Usufruto_Vitalicio.docx",
        "alt_file": "29_Contrato_Doacao_de_Bens.docx"
    },
    {
        "id": "contrato-30",
        "title": "30. Recibo de Pagamento Geral e Quitação Plena",
        "badge": "Art. 320 CC",
        "file": "30_Recibo_Pagamento_Geral_Quitacao_Plena_Irrevogavel.docx",
        "alt_file": "30_Recibo_Geral_Quitacao_Plena.docx"
    }
]

# 1. Ensure copy to alternative filenames so either link works
print("Sincronizando arquivos .docx...")
for item in CONTRACT_MAP:
    src = os.path.join(DOWNLOADS_DIR, item["file"])
    alt = os.path.join(DOWNLOADS_DIR, item["alt_file"])
    if os.path.exists(src) and src != alt:
        shutil.copyfile(src, alt)
        print(f"  Copiado: {item['file']} -> {item['alt_file']}")

# 2. Extract full text from docx to build full, dense content for web portal
def extract_docx_text(docx_path):
    doc = docx.Document(docx_path)
    lines = []
    # Check tables (like orientacao)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    lines.append(f"💡 {text}\n")
    # Paragraphs
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            lines.append(t)
    return "\n\n".join(lines)

# 3. Build updated json
updated_json = [
    {
        "id": "instrucoes",
        "title": "Instruções Iniciais e Como Usar as Minutas",
        "badge": "Comece Aqui",
        "downloadFile": "Acervo_30_Contratos_Andrade_Florio.zip",
        "downloadLabel": "Baixar Todos os 30 Contratos (.zip)",
        "content": (
            "Seja bem-vindo ao Acervo Definitivo de Contratos Imobiliários, Comerciais e Civis da banca Andrade & Florio Sociedade de Advogados.\n\n"
            "Este acervo completo contém 30 documentos essenciais redigidos com densidade jurídica integral, cláusulas protetivas expressas, fundamentação legal atualizada e notas técnicas da banca:\n\n"
            "• MÓDULO 1: Transações e Locações Imobiliárias (11 documentos)\n"
            "  Contratos de compra e venda com arras confirmatórias, locações residenciais e comerciais, comodato com cláusula anti-usucapião, laudo cautelar de vistoria, aditivos, distratos, notificações extrajudiciais, recibo de arras, corretagem e empreitada.\n\n"
            "• MÓDULO 2: Negócios e Prestação de Serviços (9 documentos)\n"
            "  Contratos de prestação de serviços com escopo blindado, contratação PJ/Freelancer sem subordinação ou pessoalidade (blindagem trabalhista), Acordo de Confidencialidade (NDA) com cláusula penal rígida, parcerias comerciais, representação autônoma, distratos e cessão de direitos.\n\n"
            "• MÓDULO 3: Vida Civil, Veículos e Finanças (10 documentos)\n"
            "  Compra e venda de veículos com comunicação de venda obrigatória (Art. 134 CTB), Confissão de Dívida com eficácia de Título Executivo Extrajudicial (Art. 784, III do CPC) para execução direta de penhora, Mútuo Financeiro, Termo de Transação e Acordo com quitação rasa, notificações de constituição em mora, procurações de gestão e venda de imóveis, declaração de gratuidade de justiça e doação com usufruto vitalício.\n\n"
            "COMO UTILIZAR:\n"
            "1. Navegue pelo menu lateral para visualizar a íntegra das cláusulas de cada contrato na tela.\n"
            "2. Clique no botão dourado 'Baixar Minuta em Word (.docx)' no topo de qualquer contrato para baixar o arquivo editável formatado em fontes jurídicas e pronto para preenchimento dos campos entre colchetes [ ].\n"
            "3. Para obter a pasta completa contendo todos os 30 contratos organizados, utilize o botão 'Baixar Acervo Completo em ZIP'."
        )
    }
]

print("Extraindo texto completo de cada contrato para o portal...")
for item in CONTRACT_MAP:
    fpath = os.path.join(DOWNLOADS_DIR, item["file"])
    if not os.path.exists(fpath):
        fpath = os.path.join(DOWNLOADS_DIR, item["alt_file"])
    
    if os.path.exists(fpath):
        full_content = extract_docx_text(fpath)
    else:
        full_content = f"Minuta do {item['title']}. Documento disponível para download."

    updated_json.append({
        "id": item["id"],
        "title": item["title"],
        "badge": item["badge"],
        "downloadFile": item["alt_file"], # Use safe canonical name
        "downloadLabel": "Baixar Minuta em Word (.docx)",
        "content": full_content
    })
    print(f"  OK: {item['id']} ({len(full_content)} caracteres de texto jurídico)")

with open(JSON_PATH, "w", encoding="utf-8") as f:
    json.dump(updated_json, f, ensure_ascii=False, indent=2)

print(f"\nArquivo {JSON_PATH} atualizado com sucesso!")

# 4. Generate master ZIP with all 30 robust .docx contracts
zip_filename = "Acervo_30_Contratos_Andrade_Florio.zip"
zip_path = os.path.join(DOWNLOADS_DIR, zip_filename)

with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
    for item in CONTRACT_MAP:
        src_path = os.path.join(DOWNLOADS_DIR, item["alt_file"])
        if os.path.exists(src_path):
            zipf.write(src_path, arcname=item["alt_file"])

print(f"ZIP Mestre criado em: {zip_path} (Tamanho: {os.path.getsize(zip_path)} bytes)")

# Also mirror zip to other convenient download paths
copies = [
    os.path.join(ROOT, "public", zip_filename),
    os.path.join(ROOT, "public", "downloads", "Acervo_Contratos_Andrade_Florio.zip"),
    os.path.join(os.path.expanduser("~"), "Downloads", zip_filename),
    os.path.join(os.path.expanduser("~"), "Desktop", zip_filename)
]

for cp in copies:
    try:
        shutil.copyfile(zip_path, cp)
        print(f"  Espelhado em: {cp}")
    except Exception as e:
        print(f"  Aviso: Não foi possível copiar para {cp}: {e}")

print("\nProcesso de sincronização 100% concluído!")
